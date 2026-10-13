"""
=============================================================================
PRICE LOGGER (REST Snapshot)  —  Logs prices for AAPL, TSLA, NVDA to Excel
=============================================================================
Uses REST snapshot polling (NOT WebSocket) so it can run alongside
ibkr_broker.py which already occupies the single WebSocket connection.

Polls /iserver/marketdata/snapshot every 1 second.

Output files (saved in same directory as this script):
    prices_AAPL_DD_MM_YY.xlsx
    prices_TSLA_DD_MM_YY.xlsx
    prices_NVDA_DD_MM_YY.xlsx

Usage:
    cd ~/Desktop/System_Trades/MACD_ENTRY_EXIT
    python3 price_logger.py          (while ibkr_broker.py is running)
=============================================================================
"""

import json
import logging
import os
import time
from datetime import datetime, timedelta
from typing import Dict, Optional

import requests
import urllib3
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

# ─── Configuration ────────────────────────────────────────────────────────────
BASE_URL    = "https://localhost:5000/v1/api"
ACCOUNT_ID  = "DUXXXXXXX"
VERIFY_SSL  = False
LOG_SYMBOLS = ["AAPL", "TSLA", "NVDA"]
LOG_DURATION_MINUTES =60         # change to 5 for quick test
POLL_INTERVAL = 1                  # seconds between polls
SESSION_FILE = os.path.join(os.path.expanduser("~"), ".ibkr_session.json")
SCRIPT_DIR   = os.path.dirname(os.path.abspath(__file__))

MD_FIELDS = ["31", "84", "85", "86", "88", "7059", "70", "71", "7762", "87"]

# ─── Logger ───────────────────────────────────────────────────────────────────
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [PRICE_LOG] %(levelname)s - %(message)s",
)
logger = logging.getLogger("price_logger")


# =============================================================================
# PRICE DATA STORE
# =============================================================================

class PriceStore:
    def __init__(self):
        self._data: Dict[str, list] = {sym: [] for sym in LOG_SYMBOLS}

    def add(self, symbol: str, row: dict):
        if symbol in self._data:
            self._data[symbol].append(row)

    def get_all(self, symbol: str) -> list:
        return list(self._data.get(symbol, []))

    def count(self, symbol: str) -> int:
        return len(self._data.get(symbol, []))

    def total(self) -> int:
        return sum(len(v) for v in self._data.values())


# =============================================================================
# IBKR SESSION  —  read-only, shares cookies with ibkr_broker.py
# =============================================================================

class SimpleSession:
    def __init__(self):
        self._session = requests.Session()
        self._session.verify = VERIFY_SSL
        self._load_cookies()

    def _load_cookies(self):
        if not os.path.exists(SESSION_FILE):
            logger.warning("No session file at %s", SESSION_FILE)
            return
        try:
            with open(SESSION_FILE) as f:
                cookies = json.load(f)
            for name, info in cookies.items():
                self._session.cookies.set(
                    name, info["value"],
                    domain=info.get("domain", "localhost"),
                    path=info.get("path", "/"),
                )
            logger.info("Loaded session cookies from %s", SESSION_FILE)
        except Exception as exc:
            logger.warning("Cookie load failed: %s", exc)

    def get(self, path, params=None):
        resp = self._session.get(f"{BASE_URL}/{path.lstrip('/')}", params=params, timeout=10)
        resp.raise_for_status()
        return resp.json()

    def post(self, path, payload=None):
        resp = self._session.post(
            f"{BASE_URL}/{path.lstrip('/')}",
            json=payload,
            headers={"Content-Type": "application/json"},
            timeout=10,
        )
        resp.raise_for_status()
        return resp.json()

    def auth_status(self):
        return self.get("/iserver/auth/status")

    def search_conid(self, symbol: str) -> Optional[int]:
        try:
            results = self.post("/iserver/secdef/search",
                                {"symbol": symbol, "name": False, "secType": "STK"})
            if results:
                return int(results[0]["conid"])
        except Exception as exc:
            logger.error("conid lookup failed for %s: %s", symbol, exc)
        return None

    def snapshot(self, conids_str: str, fields_str: str):
        return self.get("/iserver/marketdata/snapshot",
                        params={"conids": conids_str, "fields": fields_str})


# =============================================================================
# VOLUME PARSER
# =============================================================================

def _parse_vol(v) -> float:
    if v is None:
        return 0.0
    s = str(v).strip().upper()
    if s in ("", "0"):
        return 0.0
    try:
        if s.endswith("B"): return float(s[:-1]) * 1e9
        if s.endswith("M"): return float(s[:-1]) * 1e6
        if s.endswith("K"): return float(s[:-1]) * 1e3
        return float(s)
    except (ValueError, TypeError):
        return 0.0


def _parse_price(raw) -> float:
    """Parse IBKR price field — sometimes prefixed with C or H."""
    if raw is None:
        return 0.0
    try:
        return float(str(raw).lstrip("CHch"))
    except (ValueError, TypeError):
        return 0.0


# =============================================================================
# EXCEL WRITER
# =============================================================================

def save_to_excel(store: PriceStore):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    tag = datetime.now().strftime("%d_%m_%y")
    headers = [
        "Timestamp", "Last Price", "Bid", "Ask", "Spread",
        "High", "Low", "Last Size", "Bar Volume", "Daily Volume",
    ]
    col_widths = [24, 14, 14, 14, 10, 14, 14, 12, 14, 16]

    header_font  = Font(name="Arial", bold=True, color="FFFFFF", size=11)
    header_fill  = PatternFill("solid", fgColor="2F5496")
    header_align = Alignment(horizontal="center", vertical="center")
    data_font    = Font(name="Arial", size=10)
    thin_border  = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )
    alt_fill = PatternFill("solid", fgColor="F2F7FB")
    saved_files = []

    for symbol in LOG_SYMBOLS:
        rows = store.get_all(symbol)
        if not rows:
            logger.warning("[%s] No data collected — skipping Excel", symbol)
            continue

        wb = Workbook()
        ws = wb.active
        ws.title = f"{symbol} Prices"

        ws.merge_cells("A1:J1")
        title_cell = ws["A1"]
        title_cell.value = (f"{symbol} — Price Log (REST Snapshot)  |  "
                            f"{datetime.now().strftime('%B %d, %Y')}  |  "
                            f"{len(rows)} ticks")
        title_cell.font = Font(name="Arial", bold=True, size=13, color="2F5496")
        title_cell.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 30

        prices = [r["last_price"] for r in rows]
        ws.merge_cells("A2:J2")
        summary = ws["A2"]
        summary.value = (
            f"Open: {prices[0]:.4f}  |  Close: {prices[-1]:.4f}  |  "
            f"High: {max(prices):.4f}  |  Low: {min(prices):.4f}  |  "
            f"Range: {max(prices) - min(prices):.4f}  |  "
            f"Start: {rows[0]['timestamp']}  |  End: {rows[-1]['timestamp']}"
        )
        summary.font = Font(name="Arial", size=10, color="666666")
        summary.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 22
        ws.row_dimensions[3].height = 8

        header_row = 4
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=header_row, column=col_idx, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = thin_border
        ws.row_dimensions[header_row].height = 26

        for row_idx, row_data in enumerate(rows, header_row + 1):
            values = [
                row_data["timestamp"], row_data["last_price"],
                row_data["bid"], row_data["ask"], row_data["spread"],
                row_data["high"], row_data["low"], row_data["last_size"],
                row_data["bar_volume"], row_data["daily_volume"],
            ]
            use_alt = (row_idx - header_row) % 2 == 0
            for col_idx, val in enumerate(values, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = data_font
                cell.border = thin_border
                if use_alt:
                    cell.fill = alt_fill
                if col_idx in (2, 3, 4, 6, 7):
                    cell.number_format = '#,##0.0000'
                elif col_idx == 5:
                    cell.number_format = '#,##0.0000'
                elif col_idx in (8, 9, 10):
                    cell.number_format = '#,##0'

        for i, w in enumerate(col_widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = f"A{header_row + 1}"
        ws.auto_filter.ref = f"A{header_row}:J{header_row + len(rows)}"

        filename = f"prices_{symbol}_{tag}.xlsx"
        filepath = os.path.join(SCRIPT_DIR, filename)
        wb.save(filepath)
        saved_files.append(filepath)
        logger.info("[%s] Saved %d ticks → %s", symbol, len(rows), filepath)

    return saved_files


# =============================================================================
# MAIN — REST SNAPSHOT POLLING
# =============================================================================

def run_price_logger():
    print("\n" + "=" * 70)
    print("  IBKR PRICE LOGGER (REST Snapshot) — AAPL, TSLA, NVDA")
    print(f"  Polls every {POLL_INTERVAL}s for {LOG_DURATION_MINUTES} minutes")
    print(f"  Files will be saved to: {SCRIPT_DIR}")
    print("  Press Ctrl-C at any time to save and exit early")
    print("=" * 70 + "\n")

    session = SimpleSession()

    try:
        status = session.auth_status()
        if not status.get("authenticated"):
            logger.error("NOT AUTHENTICATED — is ibkr_broker.py running?")
            return
        logger.info("Authentication OK")
    except Exception as exc:
        logger.error("Cannot reach gateway: %s", exc)
        return

    # ── Resolve conids ────────────────────────────────────────────────────
    conid_map: Dict[int, str] = {}
    for sym in LOG_SYMBOLS:
        conid = session.search_conid(sym)
        if conid:
            conid_map[conid] = sym
            logger.info("Resolved %s → conid=%d", sym, conid)
        else:
            logger.error("Could not resolve %s — skipping", sym)

    if not conid_map:
        logger.error("No symbols resolved — exiting")
        return

    conids_str = ",".join(str(c) for c in conid_map.keys())
    fields_str = ",".join(MD_FIELDS)

    # Pre-flight snapshots (first 2-3 calls return empty, this warms it up)
    logger.info("Warming up snapshot stream...")
    for _ in range(3):
        try:
            session.snapshot(conids_str, fields_str)
        except Exception:
            pass
        time.sleep(1)

    # ── Poll loop ─────────────────────────────────────────────────────────
    store = PriceStore()
    start_time = datetime.now()
    end_time = start_time + timedelta(minutes=LOG_DURATION_MINUTES)
    last_prices: Dict[str, float] = {}
    last_status_min = -1

    logger.info("Polling started at %s — runs until %s",
                start_time.strftime("%H:%M:%S"), end_time.strftime("%H:%M:%S"))

    try:
        while datetime.now() < end_time:
            now = datetime.now()

            remaining = (end_time - now).total_seconds()
            mins_left = int(remaining // 60)
            if mins_left != last_status_min:
                last_status_min = mins_left
                sym_counts = "  ".join(f"{s}={store.count(s)}" for s in LOG_SYMBOLS)
                logger.info("⏱  %d min remaining  |  Total: %d ticks  |  %s",
                            mins_left, store.total(), sym_counts)

            try:
                results = session.snapshot(conids_str, fields_str)
            except Exception:
                time.sleep(POLL_INTERVAL)
                continue

            if not isinstance(results, list):
                time.sleep(POLL_INTERVAL)
                continue

            for data in results:
                conid = data.get("conid")
                if conid is None or conid not in conid_map:
                    continue

                price = _parse_price(data.get("31"))
                if price == 0:
                    continue

                symbol = conid_map[conid]

                if last_prices.get(symbol) == price:
                    continue
                last_prices[symbol] = price

                bid       = _parse_price(data.get("84"))
                ask       = _parse_price(data.get("86"))
                high      = _parse_price(data.get("70"))
                low       = _parse_price(data.get("71"))
                last_size = _parse_vol(data.get("7059", 0))
                bar_vol   = _parse_vol(data.get("7762", 0))
                daily_vol = _parse_vol(data.get("87", 0))

                row = {
                    "timestamp"   : now.strftime("%Y-%m-%d %H:%M:%S.%f")[:-3],
                    "last_price"  : price,
                    "bid"         : bid,
                    "ask"         : ask,
                    "spread"      : round(ask - bid, 4) if bid > 0 and ask > 0 else 0,
                    "high"        : high,
                    "low"         : low,
                    "last_size"   : last_size,
                    "bar_volume"  : bar_vol,
                    "daily_volume": daily_vol,
                }
                store.add(symbol, row)

                count = store.count(symbol)
                if count == 1:
                    logger.info("[%s] ✓ First tick: %.2f  bid=%.2f  ask=%.2f",
                                symbol, price, bid, ask)
                elif count % 100 == 0:
                    logger.info("[%s] %d ticks  last=%.2f", symbol, count, price)

            time.sleep(POLL_INTERVAL)

        logger.info("Timer complete — saving Excel files now...")

    except KeyboardInterrupt:
        logger.info("Interrupted by user — saving collected data now...")

    saved = save_to_excel(store)

    print("\n" + "=" * 70)
    print("  PRICE LOGGING COMPLETE")
    print(f"  Total ticks: {store.total()}")
    for s in LOG_SYMBOLS:
        print(f"    {s}: {store.count(s)} ticks")
    if saved:
        print("\n  Files saved:")
        for f in saved:
            print(f"    → {f}")
    else:
        print("\n  ⚠ No files saved (no data collected)")
    print("=" * 70 + "\n")


if __name__ == "__main__":
    run_price_logger()