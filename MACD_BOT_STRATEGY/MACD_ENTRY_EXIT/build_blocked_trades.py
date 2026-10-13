#!/usr/bin/env python3
"""
build_blocked_trades.py — list every signal the ML gate BLOCKED (skipped) today,
into BLOCKED_TRADE_<YYYY-MM-DD>.xlsx (Stock, Entry Time, Model prob_yes, ENTRY).
ENTRY defaults to "No"; edit to "Yes" for a block you'd actually have taken.
Preserves your edits on every re-run (cron-safe). Optional arg YYYY-MM-DD.
"""
import re, datetime, sys, os
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side as BSide
from openpyxl.worksheet.datavalidation import DataValidation

BASE  = "/home/azureuser/MACD_ENTRY_EXIT"
today = datetime.datetime.strptime(sys.argv[1], "%Y-%m-%d") if len(sys.argv) > 1 else datetime.datetime.now()
L     = f"{BASE}/logs/tws_{today.strftime('%d_%m_%y')}.txt"
OUT   = f"{BASE}/BLOCKED_TRADE_{today.strftime('%Y-%m-%d')}.xlsx"

pat = re.compile(r"(\d\d:\d\d:\d\d),\d+ .*\[(\w+)\] ML GATE .* SKIP .*prob_yes=([\d.]+)")
rows = []
try:
    for line in open(L):
        m = pat.search(line)
        if m:
            rows.append((m.group(2), m.group(1), float(m.group(3))))   # (sym, time, prob)
except FileNotFoundError:
    print(f"no log {L}"); raise SystemExit(0)

# preserve existing ENTRY labels keyed by (symbol, time)
existing = {}
if os.path.exists(OUT):
    try:
        _ws = load_workbook(OUT).active
        for r in range(2, _ws.max_row + 1):
            s = _ws.cell(row=r, column=1).value; t = _ws.cell(row=r, column=2).value
            if s and t:
                existing[(str(s), str(t))] = _ws.cell(row=r, column=4).value
    except Exception:
        pass

wb = Workbook(); ws = wb.active; ws.title = "Blocked " + today.strftime("%Y-%m-%d")
headers = ["Stock", "Entry Time", "Model prob_yes", "ENTRY"]
hf = PatternFill("solid", fgColor="7A2E2E"); hfont = Font(bold=True, color="FFFFFF", size=11)
thin = BSide(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal="center", vertical="center")
for c, h in enumerate(headers, 1):
    cell = ws.cell(row=1, column=c, value=h); cell.fill = hf; cell.font = hfont; cell.alignment = ctr; cell.border = bd
for i, (sym, t, prob) in enumerate(rows, start=2):
    lbl = existing.get((sym, t))
    lbl = lbl if lbl not in (None, "") else "No"      # default N (=No), user edits to Yes
    vals = [sym, t, round(prob, 4), lbl]
    for c, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=c, value=v); cell.alignment = ctr; cell.border = bd
if rows:
    dv = DataValidation(type="list", formula1='"Yes,No,Maybe"', allow_blank=True)
    dv.add(f"D2:D{len(rows)+1}"); ws.add_data_validation(dv)
for col, w in {"A": 10, "B": 12, "C": 15, "D": 10}.items():
    ws.column_dimensions[col].width = w
ws.freeze_panes = "A2"
wb.save(OUT)
print(f"Wrote {OUT} with {len(rows)} blocked signals")
