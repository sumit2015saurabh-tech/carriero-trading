#!/usr/bin/env python3
"""
add_portfolios_sheet.py
=======================
One-shot migration: adds a "Portfolios" sheet and a "Portfolio" column to
strategy_config.xlsx so per-portfolio defaults can be defined.

What it does:
  1. Makes a timestamped backup of the existing xlsx.
  2. If a "Portfolios" sheet already exists → leaves it alone.
  3. Otherwise creates a "Portfolios" sheet with:
       - Title row, header row, one "DEFAULT" portfolio row pre-populated
         from the current Global Settings so the inheritance chain is a
         no-op until the user customises it.
  4. If the "Stocks" sheet lacks a "Portfolio" column (column 22 / V) →
     adds it as the last column, header "Portfolio". No existing data is
     touched. Existing stocks default to DEFAULT (cell left blank).

Safe to re-run — idempotent: if the sheet/column already exist, no-op.

Usage:
    python add_portfolios_sheet.py
    python add_portfolios_sheet.py /path/to/strategy_config.xlsx
"""
from __future__ import annotations

import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

# ── Paths ────────────────────────────────────────────────────────────────────

if len(sys.argv) > 1:
    XLSX = Path(sys.argv[1])
else:
    XLSX = Path(__file__).parent / "strategy_config.xlsx"

if not XLSX.exists():
    print(f"ERROR: {XLSX} not found.")
    sys.exit(1)

# ── 1. Backup ────────────────────────────────────────────────────────────────

ts  = datetime.now().strftime("%H%M_%d%m")
bak = XLSX.with_suffix(f".xlsx.bak.{ts}_pre_portfolios")
shutil.copy2(str(XLSX), str(bak))
print(f"Backup created: {bak.name}")

# ── 2. Load workbook ─────────────────────────────────────────────────────────

wb = openpyxl.load_workbook(str(XLSX))
print(f"Existing sheets: {wb.sheetnames}")

# ── 3. Read Global Settings to seed the DEFAULT portfolio ────────────────────

defaults = {
    "macd_resolution"   : "30s",
    "macd_fast"         : 12,
    "macd_slow"         : 26,
    "macd_signal_period": 9,
    "macd_price_source" : "close",
    "mpl"               : 0.6,
    "signal_l"          : 0,
    "signal_s"          : 0,
    "mean_multiplier_l" : 1.0,
    "mean_multiplier_s" : 1.0,
    "macdle_blw_zline"  : "N",
    "macdse_abv_zline"  : "N",
    "bem"               : 1.0,
    "trailing_perc"     : 0.1,
}
try:
    gs = wb["Global Settings"]
    for row in gs.iter_rows(min_row=3, values_only=True):
        if not row or not row[0]:
            continue
        key = str(row[0]).strip()
        val = row[1] if len(row) > 1 else None
        if val is None:
            continue
        mapping = {
            "MACD_Resolution"  : "macd_resolution",
            "MACD_Fast"        : "macd_fast",
            "MACD_Slow"        : "macd_slow",
            "MACD_Signal"      : "macd_signal_period",
            "MACD_PriceSource" : "macd_price_source",
            "MPL_dollar_Amount": "mpl",
            "Signal_L"         : "signal_l",
            "Signal_S"         : "signal_s",
            "MeanMultiplier_L" : "mean_multiplier_l",
            "MeanMultiplier_S" : "mean_multiplier_s",
            "MACDLE_BLW_ZLine" : "macdle_blw_zline",
            "MACDSE_ABV_ZLine" : "macdse_abv_zline",
            "BEMultiplier"     : "bem",
            "TrailingPerc"     : "trailing_perc",
        }
        if key in mapping:
            defaults[mapping[key]] = val
    print(f"Seeded DEFAULT portfolio from Global Settings")
except KeyError:
    print("WARN: Global Settings sheet not found — using built-in defaults")

# ── 4. Add Portfolios sheet (if missing) ─────────────────────────────────────

if "Portfolios" in wb.sheetnames:
    print("'Portfolios' sheet already exists — skipping create.")
else:
    ws = wb.create_sheet("Portfolios")

    title_font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_fill = PatternFill("solid", fgColor="2E5266")
    hdr_font   = Font(name="Calibri", size=11, bold=True)
    hdr_fill   = PatternFill("solid", fgColor="DDE6E8")

    ws["A1"] = ("📂  PORTFOLIO LIST  —  Each portfolio defines DEFAULTS for stocks "
                "assigned to it. Inheritance: per-stock field → portfolio → global.")
    ws["A1"].font = title_font
    ws["A1"].fill = title_fill
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 22
    ws.merge_cells("A1:P1")

    headers = [
        "Portfolio Name", "MACD Resolution",
        "MACD Fast", "MACD Slow", "MACD Signal", "Price Source",
        "MPL", "Signal_L", "Signal_S",
        "Mean Multiplier_L", "Mean Multiplier_S",
        "MACDLE_BLW_Zline", "MACDSE_ABV_Zline",
        "BEM", "TrailingPerc", "Notes",
    ]
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=c, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    # Seed DEFAULT row
    default_row = [
        "DEFAULT",
        defaults["macd_resolution"],
        defaults["macd_fast"], defaults["macd_slow"], defaults["macd_signal_period"],
        defaults["macd_price_source"],
        defaults["mpl"],
        defaults["signal_l"], defaults["signal_s"],
        defaults["mean_multiplier_l"], defaults["mean_multiplier_s"],
        defaults["macdle_blw_zline"], defaults["macdse_abv_zline"],
        defaults["bem"], defaults["trailing_perc"],
        "Default portfolio — values mirror Global Settings.",
    ]
    for c, v in enumerate(default_row, start=1):
        ws.cell(row=3, column=c, value=v)

    # Add a couple of suggested-empty example rows
    ws.cell(row=4, column=1, value="TECH_30S")
    ws.cell(row=4, column=2, value="30s")
    ws.cell(row=4, column=16, value="Example: 30-sec-bar tech stocks. Edit values then assign in Stocks sheet.")
    ws.cell(row=5, column=1, value="SEMIS_1MIN")
    ws.cell(row=5, column=2, value="1min")
    ws.cell(row=5, column=16, value="Example: 1-min semiconductor names.")

    # Column widths
    widths = [16, 16, 11, 11, 12, 14, 8, 10, 10, 18, 18, 17, 17, 8, 13, 50]
    for c, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + c)].width = w

    ws.freeze_panes = "A3"
    print("Created 'Portfolios' sheet with DEFAULT row + 2 examples.")

# ── 5. Add Portfolio column to Stocks (if missing) ──────────────────────────

ws_s = wb["Stocks"]
# Inspect headers row 2; "Portfolio" header is what we look for.
existing_headers = [
    str(ws_s.cell(2, c).value or "").strip().lower()
    for c in range(1, ws_s.max_column + 1)
]
if "portfolio" in existing_headers:
    print(f"'Portfolio' column already present in Stocks (col {existing_headers.index('portfolio')+1}).")
else:
    new_col = ws_s.max_column + 1
    hdr = ws_s.cell(row=2, column=new_col, value="Portfolio")
    # Match header style from neighbours (col 21 is the rightmost existing header)
    if new_col >= 2:
        ref = ws_s.cell(row=2, column=new_col - 1)
        if ref.font is not None:
            hdr.font = Font(name=ref.font.name, size=ref.font.size, bold=ref.font.bold)
        if ref.fill is not None and ref.fill.patternType:
            hdr.fill = PatternFill("solid", fgColor=ref.fill.fgColor.rgb if ref.fill.fgColor else "DDE6E8")
        hdr.alignment = Alignment(horizontal="center", vertical="center")
    ws_s.column_dimensions[hdr.column_letter].width = 14
    print(f"Added 'Portfolio' column to Stocks at position {new_col}.")

# ── 6. Save ──────────────────────────────────────────────────────────────────

wb.save(str(XLSX))
print(f"Saved updated workbook → {XLSX.name}")
print()
print("Next steps:")
print("  1. Open strategy_config.xlsx in Excel.")
print("  2. Fill in the 'Portfolio' column on the Stocks sheet (leave blank = DEFAULT).")
print("  3. Customize the DEFAULT row on the Portfolios sheet — these values are the")
print("     fallback for any stock that leaves a field blank.")
print("  4. Add more portfolio rows if you want grouped templates (e.g. TECH_30S).")
print("  5. Save the file. excel_sync.py will pick it up automatically.")
