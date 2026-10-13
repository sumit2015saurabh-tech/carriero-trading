import csv, json
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side as BSide

BASE = "/home/azureuser/MACD_ENTRY_EXIT"
CFG = json.load(open(f"{BASE}/config.json"))["Stocks"]
K_ATR = 3.0
OUT = f"{BASE}/trades_chandelier_7d.xlsx"

def ema(series, n):
    k = 2.0/(n+1); out = []; e = None
    for x in series:
        e = x if e is None else x*k + e*(1-k); out.append(e)
    return out

def cfgf(sym, key, d):
    try: return float(CFG.get(sym, {}).get(key, d))
    except Exception: return d

def fmt_dt(dt):
    p = dt.split()
    d = p[0]; t = p[1] if len(p) > 1 else ""
    return f"{d[:4]}-{d[4:6]}-{d[6:8]}", t

def run_symbol(sym, bars):
    bars.sort(key=lambda x: x[1])
    dts = [b[1] for b in bars]
    closes = [b[5] for b in bars]; highs = [b[3] for b in bars]; lows = [b[4] for b in bars]; days = [b[0] for b in bars]
    n = len(bars)
    if n < 30: return []
    e12 = ema(closes, 12); e26 = ema(closes, 26)
    macd = [e12[i]-e26[i] for i in range(n)]; signal = ema(macd, 9)
    tr = [highs[0]-lows[0]] + [max(highs[i]-lows[i], abs(highs[i]-closes[i-1]), abs(lows[i]-closes[i-1])) for i in range(1, n)]
    atr = [sum(tr[max(0, i-13):i+1])/len(tr[max(0, i-13):i+1]) for i in range(n)]
    mpl = cfgf(sym, "mpl", 0.8); qty = cfgf(sym, "quantity", 500)
    mml = cfgf(sym, "mean_multiplier_l", 1.0); mms = cfgf(sym, "mean_multiplier_s", 1.0)
    sneg = spos = 0.0; ncnt = pcnt = 0; curday = None
    pos = 0; entry = 0.0; stop = 0.0; peak = 0.0; ei = 0
    trades = []

    def book(xi, xpx, reason):
        nonlocal pos
        side = "LONG" if pos > 0 else "SHORT"
        pnl = (xpx - entry) * (1 if pos > 0 else -1) * qty
        d, et = fmt_dt(dts[ei]); _, xt = fmt_dt(dts[xi])
        trades.append([d, side, et, round(entry, 2), xt, round(xpx, 2), round(pnl, 0), reason, dts[ei]])
        pos = 0

    for i in range(n):
        if days[i] != curday:
            if pos != 0: book(i-1, closes[i-1], "End of day flat")
            curday = days[i]; sneg = spos = 0.0; ncnt = pcnt = 0
        s = signal[i]
        if s < 0: sneg = (sneg*ncnt + s)/(ncnt+1); ncnt += 1
        elif s > 0: spos = (spos*pcnt + s)/(pcnt+1); pcnt += 1
        if i < 2: continue
        if pos == 0:
            tu = macd[i] > macd[i-1]; below = macd[i] < signal[i]
            td = macd[i] < macd[i-1]; above = macd[i] > signal[i]
            xL = (macd[i-1] <= signal[i-1]) and (macd[i] > signal[i])
            xS = (macd[i-1] >= signal[i-1]) and (macd[i] < signal[i])
            gl = (sneg == 0.0) or (sneg >= mml*signal[i]); gs = (spos == 0.0) or (spos <= mms*signal[i])
            if (tu and below) or (xL and gl):
                pos = 1; entry = closes[i]; peak = closes[i]; stop = entry - mpl; ei = i
            elif (td and above) or (xS and gs):
                pos = -1; entry = closes[i]; peak = closes[i]; stop = entry + mpl; ei = i
            continue
        if pos > 0:
            peak = max(peak, highs[i])
            if lows[i] <= stop:
                book(i, stop, "ATR trail stop (profit)" if stop > entry else "Stop loss"); continue
            stop = max(stop, peak - K_ATR*atr[i])
        else:
            peak = min(peak, lows[i])
            if highs[i] >= stop:
                book(i, stop, "ATR trail stop (profit)" if stop < entry else "Stop loss"); continue
            stop = min(stop, peak + K_ATR*atr[i])
    if pos != 0: book(n-1, closes[-1], "End of day flat")
    return trades

rows = defaultdict(list)
for r in csv.DictReader(open(f"{BASE}/hist_7d_3min.csv")):
    dt = r["datetime"].strip()
    try: o, h, l, c = float(r["open"]), float(r["high"]), float(r["low"]), float(r["close"])
    except ValueError: continue
    rows[r["symbol"]].append((dt.split()[0], dt, o, h, l, c))

wb = Workbook(); wb.remove(wb.active)
hdr = ["Date", "Side", "Entry Time", "Entry Price", "Exit Time", "Exit Price", "P&L ($)", "Exit Reason"]
hf = PatternFill("solid", fgColor="1F4E78"); hfont = Font(bold=True, color="FFFFFF")
thin = BSide(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal="center", vertical="center")
grand = 0.0; counts = {}
for sym in sorted(rows):
    tr = run_symbol(sym, rows[sym])
    tr.sort(key=lambda x: x[8])     # sort by entry datetime
    ws = wb.create_sheet(sym[:31])
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(row=1, column=c, value=h); cell.fill = hf; cell.font = hfont; cell.alignment = ctr; cell.border = bd
    tot = 0.0
    for ri, t in enumerate(tr, start=2):
        for c, v in enumerate(t[:8], 1):
            cell = ws.cell(row=ri, column=c, value=v); cell.alignment = ctr; cell.border = bd
        tot += t[6]
    last = len(tr) + 2
    ws.cell(row=last, column=6, value="TOTAL").font = Font(bold=True)
    ws.cell(row=last, column=7, value=round(tot, 0)).font = Font(bold=True)
    for col, w in {"A":12,"B":7,"C":11,"D":12,"E":11,"F":11,"G":11,"H":22}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    grand += tot; counts[sym] = (len(tr), round(tot, 0))
wb.save(OUT)
print(f"Saved {OUT}")
print(f"sheets: {len(counts)} | total trades: {sum(c[0] for c in counts.values())} | grand P&L: {grand:+.0f}")
