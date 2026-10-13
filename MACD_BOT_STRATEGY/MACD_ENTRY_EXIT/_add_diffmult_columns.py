"""
Add DiffMultiplier_L and DiffMultiplier_S columns to the Stocks and Indexes
sheets of strategy_config.xlsx. Default value: 1.0 (matches the prior global
default in entry_logic.py).

Idempotent — checks for an existing DiffMultiplier_L header before adding.
"""
import sys, shutil
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

XL_PATH = sys.argv[1] if len(sys.argv) > 1 else "strategy_config.xlsx"

DEFAULT_DIFF_MULT = 1.0

shutil.copy(XL_PATH, XL_PATH + ".bak.before_diffmult")
wb = load_workbook(XL_PATH)

def add_pair(ws, sheet_name: str):
    cur_max = ws.max_column
    # Check if already present
    for c in range(1, cur_max + 1):
        h = (ws.cell(row=2, column=c).value or "")
        if "diffmultiplier_l" in str(h).lower().replace(" ", ""):
            print(f"[{sheet_name}] DiffMultiplier already present at col {c} — skipping")
            return
    c_l = cur_max + 1
    c_s = cur_max + 2
    ws.cell(row=1, column=c_l, value="Diff")
    ws.cell(row=1, column=c_s, value="Diff")
    ws.cell(row=2, column=c_l, value="DiffMultiplier_L")
    ws.cell(row=2, column=c_s, value="DiffMultiplier_S")
    for r in range(3, ws.max_row + 1):
        if ws.cell(row=r, column=1).value:
            if ws.cell(row=r, column=c_l).value is None:
                ws.cell(row=r, column=c_l, value=DEFAULT_DIFF_MULT)
            if ws.cell(row=r, column=c_s).value is None:
                ws.cell(row=r, column=c_s, value=DEFAULT_DIFF_MULT)
    for c in (c_l, c_s):
        cell = ws.cell(row=2, column=c)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    print(f"[{sheet_name}] added DiffMultiplier_L at col {c_l}, "
          f"DiffMultiplier_S at col {c_s}")

add_pair(wb["Stocks"],  "Stocks")
add_pair(wb["Indexes"], "Indexes")

wb.save(XL_PATH)
print(f"\nSaved {XL_PATH}.  Backup: {XL_PATH}.bak.before_diffmult")
