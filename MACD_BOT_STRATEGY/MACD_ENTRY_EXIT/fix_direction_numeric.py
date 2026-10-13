import openpyxl
from openpyxl.styles import Alignment
P="strategy_config.xlsx"
wb=openpyxl.load_workbook(P)
ws=wb["Stocks"]
COL=28  # AB
ws.cell(1,COL).value="Direction"
ws.cell(2,COL).value="1=S 2=L 3=both 4=off"
n=0
for r in range(3, ws.max_row+1):
    if ws.cell(r,1).value not in (None,""):
        ws.cell(r,COL).value=3          # default = both (numeric, like ExecutionType)
        ws.cell(r,COL).alignment=Alignment(horizontal="center",vertical="center")
        n+=1
wb.save(P)
print("Direction (col AB) set NUMERIC =3 (both) on", n, "named rows; header='Direction' / '1=S 2=L 3=both 4=off'")
